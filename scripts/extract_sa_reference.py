#!/usr/bin/env python3
"""
Extract the SemiAnalysis AI-Cloud TCO model inputs the benchmarks page shows as references
into store/sa_reference.json. Runs LOCALLY against the downloaded workbook (subscription
content; internal use only):

    python3 scripts/extract_sa_reference.py "~/Downloads/AI-Cloud-TCO-Model-Update-August-10-2026-SKU-1 (1).xlsx"

What is taken, per tracked tier (Neocloud Giant profile):
  - Rental Price Forecasts: the "Base Case - <SKU>" monthly market rental path ($/hr/GPU)
  - Full TCO: capex per logical GPU (r191), operating cost per GPU-month (r132), capital /
    operating / total cost per GPU-hour (r145-147), WACC (r197), useful life (r198),
    5-year calibrated pricing (r356) and realised pricing (r358)
  - Cost-Value Visualization: SemiAnalysis' own floor (price for ~15.6% project IRR on a
    5-year contract with 15% prepay) and value-based ceiling blocks where they exist
The page never pools any of this into a mark; it is a labelled model reference.
"""
import datetime as dt
import json
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "store" / "sa_reference.json"

# our tier -> (Rental Price Forecasts row label prefix, Full TCO 'Cluster Configuration' label, customer type)
TIERS = {
    "H100": ("Base Case - H100", "H100 SXM", "Neocloud Giant"),
    "H200": ("Base Case - H200", "H200 SXM", "Neocloud Giant"),
    "B200": ("Base Case - B200", "B200 SXM 8xHBM 1000W", "Neocloud Giant"),
    "GB200": ("Base Case - GB200 NVL72", "GB200 NVL72", "Neocloud Giant"),
    "B300": ("Base Case - B300", "B300 1200W", "Neocloud Giant"),
    "GB300": ("Base Case - GB300", "GB300 NVL 72", "Neocloud Giant"),
    "VR": ("Base Case - VR NVL72", "VR NVL72 2300W", "Neocloud Giant"),
}
FULL_TCO_ROWS = {"capex_per_gpu_usd": 191, "opex_per_gpu_month_usd": 132, "capital_cost_per_hour": 145,
                 "operating_cost_per_hour": 146, "total_cost_per_hour": 147, "wacc": 197, "useful_life_years": 198,
                 "calibrated_5y_price": 356, "realised_price": 358,
                 "all_in_watts_per_gpu": 77, "gpus_per_server": 139}   # inputs of SemiAnalysis' Business Case IRR (scripts/sa_irr_floor.py)


def main(argv) -> int:
    if not argv:
        print(__doc__)
        return 2
    path = Path(argv[0]).expanduser()
    import openpyxl  # local dependency only
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    version = re.search(r"(\w+-\d{1,2}-\d{4})", path.name)
    out = {"_source": f"SemiAnalysis AI-Cloud TCO model, file {path.name}", "_extracted": dt.date.today().isoformat(),
           "_version": version.group(1) if version else None,
           "_note": "Subscription content, internal use only. Base-case market rental path is a modelled monthly price, not observed offers. "
                    "Full TCO figures are the Neocloud Giant profile. Never pooled into marks.",
           "tiers": {}}

    # ---- rental price paths
    ws = wb["Rental Price Forecasts"]
    rows = list(ws.iter_rows(min_row=130, max_row=170, values_only=True))
    hdr = next(r for r in rows if sum(isinstance(v, dt.datetime) for v in r) > 12)
    col_month = {i: hdr[i].strftime("%Y-%m") for i in range(len(hdr)) if isinstance(hdr[i], dt.datetime)}
    for tier, (label, _, _) in TIERS.items():
        r = next((r for r in rows if r[1] and str(r[1]).startswith(label)), None)
        path_ = {col_month[i]: round(float(r[i]), 4) for i in col_month if r and r[i] is not None} if r else {}
        out["tiers"][tier] = {"rental_path_monthly": path_}

    # ---- Full TCO per SKU (Neocloud Giant column)
    ws = wb["Full TCO"]
    grid = {i: r for i, r in enumerate(ws.iter_rows(min_row=1, max_row=360, values_only=True), start=1)}
    cfg, cust = grid[7], grid[8]
    def col_for(cfg_label, customer):
        for j in range(len(cfg)):
            c = str(cfg[j] or "").replace("\n", " ").strip()
            if c == cfg_label and str(cust[j] or "").strip() == customer:
                return j
        return None
    for tier, (_, cfg_label, customer) in TIERS.items():
        j = col_for(cfg_label, customer)
        t = out["tiers"][tier]
        t["full_tco_column"] = cfg_label if j is not None else None
        for k, rn in FULL_TCO_ROWS.items():
            v = grid[rn][j] if j is not None else None
            t[k] = round(float(v), 4) if isinstance(v, (int, float)) else None

    # ---- SemiAnalysis' own floor / ceiling blocks (Cost-Value Visualization)
    ws = wb["Cost-Value Visualization"]
    cv = list(ws.iter_rows(min_row=1, max_row=200, max_col=8, values_only=True))
    blocks, cur = [], None
    for r in cv:
        lab = str(r[1] or "").strip()
        if lab.startswith("Floor Price (Cost-Based"):
            cur = {"sku": prev_label, "floor_irr": None, "comparison_sku": None, "comparison_tflops": None, "comparison_5y_price": None,
                   "sku_tflops": None, "ceiling": None}
            blocks.append(cur)
        elif cur is not None:
            if lab.startswith("Floor Price for"): cur["floor_irr"] = r[3]
            elif lab == "SKU of Comparison": cur["comparison_sku"] = r[3]
            elif lab.startswith("Marketed TFLOPS") and cur["comparison_tflops"] is None: cur["comparison_tflops"] = r[3]
            elif lab.startswith("Marketed TFLOPS"): cur["sku_tflops"] = r[3]
            elif lab.startswith("Market Rental Price for 5-year"): cur["comparison_5y_price"] = r[3]
            elif lab == "Ceiling Price": cur["ceiling"] = {"unit": r[2], "value": r[3]}
        if lab and not lab.startswith(("Floor", "Ceiling", "SKU", "Marketed", "Market", "Rental", "Current", "5-Year", "Project", "Note", "Chart", "Vertical", "Horizontal", "GB300 NVL72 breakeven", "Assumes")):
            prev_label = lab
    out["semianalysis_floor_ceiling_blocks"] = blocks

    OUT.write_text(json.dumps(out, indent=1))
    for tier, t in out["tiers"].items():
        p = t["rental_path_monthly"]
        print(f"{tier:6} col={str(t.get('full_tco_column'))[:22]:22} capex={t.get('capex_per_gpu_usd')} opex/mo={t.get('opex_per_gpu_month_usd')} "
              f"cost/hr={t.get('total_cost_per_hour')} wacc={t.get('wacc')} path {min(p) if p else '–'}..{max(p) if p else '–'} now={p.get(dt.date.today().strftime('%Y-%m'))}")
    print("SA floor/ceiling blocks:", [(b["sku"], b["floor_irr"], b.get("ceiling")) for b in blocks])
    print(f"wrote {OUT}")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
